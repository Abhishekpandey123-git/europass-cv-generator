"""
Export a CV to Word (.docx) with the same structure and styling as the PDF.

Word files are meant for further editing, so the layout is close to the PDF
but not pixel-identical. Word uses the font installed on the viewer's computer;
if Open Sans is missing there, Word substitutes a similar font.
"""

from __future__ import annotations

import base64
import html
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN

from cv_design import GENERIC_SECTIONS, FONTS, SECTION_LABELS, ensure_ids, mix_with_white, normalize

# inline formatting the editor can apply
_RICH = re.compile(r"</?(b|strong|i|em|u|s|br)\s*/?>", re.I)
_ALIGN = {"left": ALIGN.LEFT, "center": ALIGN.CENTER, "right": ALIGN.RIGHT, "justify": ALIGN.JUSTIFY}


def rich_parts(text: str):
    """Split '<b>a</b> b' into [(text, {bold: True}), (' b', {})] for Word runs."""
    out, marks = [], {"b": 0, "i": 0, "u": 0, "s": 0}
    alias = {"strong": "b", "em": "i"}
    pos = 0
    for m in _RICH.finditer(str(text or "")):
        chunk = text[pos:m.start()]
        if chunk:
            out.append((html.unescape(chunk), dict(marks)))
        tag = alias.get(m.group(1).lower(), m.group(1).lower())
        if tag == "br":
            out.append(("\n", dict(marks)))
        elif m.group(0).startswith("</"):
            marks[tag] = max(0, marks[tag] - 1)
        else:
            marks[tag] += 1
        pos = m.end()
    rest = text[pos:] if isinstance(text, str) else ""
    if rest:
        out.append((html.unescape(rest), dict(marks)))
    return out or [("", {})]

GREY_DATE = RGBColor(0x75, 0x75, 0x75)
GREY_ORG = RGBColor(0x6B, 0x6B, 0x6B)
GREY_SEP = RGBColor(0xD3, 0xD3, 0xD3)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def fmt_date(value) -> str:
    if not value:
        return ""
    parts = str(value).split("-")
    return "/".join(reversed(parts))


def bottom_border(paragraph, hex_color: str, size: int = 6):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "1"), ("w:color", hex_color.lstrip("#"))):
        bottom.set(qn(k), v)
    borders.append(bottom)
    ppr.append(borders)


def cell_border_bottom(cell, hex_color="C6C6C6"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "right", "bottom"):
        el = OxmlElement(f"w:{side}")
        if side == "bottom":
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), hex_color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcpr.append(borders)


class Writer:
    """Builds the .docx. Section names, spacing and alignment follow the design."""

    def __init__(self, cv: dict):
        ensure_ids(cv)
        self.cv = cv
        self.design, self.sections = normalize(cv)
        s = self.design["scale"]
        self.size_body = Pt(9 * s)
        self.size_section = Pt(10 * s)
        self.size_name = Pt(14 * s)
        self.accent = rgb(self.design["accent"])
        self.text = rgb(self.design["text"])
        self.doc = Document()
        self.container = self.doc   # where para() adds paragraphs (the page, or a table cell)
        self._setup()

    # -------------------------------------------------------------- setup ---
    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        margin = Mm(self.design["margin_mm"])
        sec.left_margin = sec.right_margin = margin
        sec.top_margin, sec.bottom_margin = Mm(10), Mm(12)

        font_name = FONTS[self.design["font"]]["docx"]
        normal = self.doc.styles["Normal"]
        normal.font.name = font_name
        normal.font.size = self.size_body
        normal.font.color.rgb = self.text
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font_name)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(1.5)
        pf.line_spacing = max(1.0, float(self.design["line_height"]) * 0.78)

        if self.design["page_numbers"]:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self._field(p, "PAGE"); r = p.add_run(" / "); r.font.size = Pt(7.5)
            self._field(p, "NUMPAGES")

    def _field(self, paragraph, code):
        run = paragraph.add_run()
        run.font.size = Pt(7.5)
        run.font.color.rgb = GREY_DATE
        for tag, text in (("begin", None), (None, code), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
            run._r.append(el)

    # ------------------------------------------------------------ helpers ---
    def para(self, space_after=1.5):
        p = self.container.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def rich(self, p, text, color=None, size=None):
        """Add text to a paragraph, keeping bold / italic / underline marks."""
        for chunk, marks in rich_parts(text):
            for i, line in enumerate(chunk.split("\n")):
                if i:
                    p.add_run().add_break()
                if not line:
                    continue
                r = self.run(p, line, bold=bool(marks.get("b")), italic=bool(marks.get("i")),
                             color=color, size=size)
                if marks.get("u"):
                    r.underline = True
                if marks.get("s"):
                    r.font.strike = True
        return p

    def align(self, p, path):
        a = self.design["block_align"].get(path)
        if a in _ALIGN:
            p.alignment = _ALIGN[a]
        return p

    def run(self, p, text, bold=False, italic=False, color=None, size=None, semibold=False):
        r = p.add_run(text)
        r.bold = bold or semibold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
        if size is not None:
            r.font.size = size
        return r

    def L(self, key, default):
        return self.design.get("labels", {}).get(key, default)

    def labelled(self, pairs, p=None):
        p = p or self.para()
        for i, (label, value) in enumerate(pairs):
            if i:
                self.run(p, "  |  ", color=GREY_SEP)
            self.run(p, f"{label}: ", bold=True)
            self.run(p, str(value))
        return p

    def label(self, key: str) -> str:
        return self.design["section_titles"].get(key) or SECTION_LABELS[key]

    def section_title(self, title):
        p = self.para(space_after=4)
        p.paragraph_format.space_before = Pt(self.design["section_gap"] * 2.2)
        p.paragraph_format.keep_with_next = True
        self.run(p, title.upper(), bold=True, color=self.accent, size=self.size_section)
        bottom_border(p, mix_with_white(self.design["accent"], 0.45), size=6)

    def dates(self, item):
        p = self.para(space_after=0.5)
        p.paragraph_format.keep_with_next = True
        end = fmt_date(item.get("end")) if item.get("end") else "Current"
        self.run(p, f"[ {fmt_date(item.get('start'))} – {end} ]", color=GREY_DATE)

    def entry_head(self, item, title_key, org_key):
        self.dates(item)
        p = self.para(space_after=0.5); p.paragraph_format.keep_with_next = True
        self.run(p, item.get(title_key, ""), bold=True, color=self.accent)
        if item.get(org_key):
            p = self.para(space_after=0.5); p.paragraph_format.keep_with_next = True
            self.run(p, item[org_key], bold=True, italic=True, color=GREY_ORG)
        loc = [(k, item[f]) for k, f in (("City", "city"), ("Country", "country")) if item.get(f)]
        if loc:
            self.labelled(loc)

    def description(self, desc, path=None):
        if not desc:
            return
        if isinstance(desc, str):
            for line in re.split(r"<br\s*/?>|\n", desc):
                if line.strip():
                    self.align(self.rich(self.para(), line.strip()), path)
        else:
            for i, item in enumerate(desc):
                p = self.container.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                self.rich(p, item)
        self.para(space_after=3)

    # ----------------------------------------------------------- sections ---
    def header(self):
        p = self.cv.get("personal", {})
        photo = p.get("photo")
        photo_data = None
        if photo and str(photo).startswith("data:image"):
            try:
                photo_data = base64.b64decode(photo.split(",", 1)[1])
            except Exception:
                photo_data = None

        cell = None
        if photo_data:
            # photo beside the details, as in the PDF: a borderless two-column table
            usable = Mm(210 - 2 * self.design["margin_mm"])
            table = self.doc.add_table(rows=1, cols=2)
            table.autofit = False
            widths = (Mm(32), usable - Mm(32))
            for col, w in zip(table.columns, widths):
                col.width = w
            for c, w in zip(table.rows[0].cells, widths):
                c.width = w
            try:
                table.rows[0].cells[0].paragraphs[0].add_run().add_picture(io.BytesIO(photo_data), width=Mm(28))
            except Exception:
                pass
            cell = table.rows[0].cells[1]
            self.container = cell

        name = self.para(space_after=3)
        self.run(name, f"{p.get('first_name', '')} {p.get('last_name', '')}".strip(),
                 bold=True, color=self.accent, size=self.size_name)

        line1 = []
        if p.get("date_of_birth"):
            line1.append((self.L("dob", "Date of birth"), fmt_date(p["date_of_birth"])))
        if p.get("nationalities"):
            line1.append((self.L("nationality", "Nationality"), ", ".join(n.get("name", "") for n in p["nationalities"])))
        if p.get("gender"):
            line1.append((self.L("gender", "Gender"), p["gender"]))
        if line1:
            self.labelled(line1)

        line2 = []
        ph = p.get("phone") or {}
        if ph.get("number"):
            num = (f"(+{ph['country_dialing']}) " if ph.get("country_dialing") else "") + ph["number"]
            line2.append((self.L("phone", "Phone number"), num + (f" ({ph['type']})" if ph.get("type") else "")))
        if p.get("email"):
            line2.append((self.L("email", "Email address"), p["email"]))
        if p.get("website"):
            line2.append((self.L("website", "Website"), p["website"]))
        if line2:
            self.labelled(line2)

        a = p.get("address") or {}
        if a.get("street") or a.get("city") or a.get("country"):
            city = f"{a.get('postal_code', '')} {a.get('city', '')}".strip()
            value = ", ".join(x for x in (a.get("street"), city, a.get("country")) if x)
            last = self.labelled([(self.L("address", "Address"), value + (f" ({a['type']})" if a.get("type") else ""))])
        else:
            last = self.para()

        if cell is not None:
            first = cell.paragraphs[0]           # the empty paragraph every new cell starts with
            if not first.text and len(cell.paragraphs) > 1:
                first._element.getparent().remove(first._element)
            self.container = self.doc
            last = self.para(space_after=2)
        bottom_border(last, "C6C6C6", size=4)

    def about(self):
        if self.cv.get("about"):
            self.section_title(self.label("about"))
            self.align(self.rich(self.para(), self.cv["about"]), "about")

    def work_experience(self):
        jobs = self.cv.get("work_experience") or []
        if jobs:
            self.section_title(self.label("work_experience"))
            for i, job in enumerate(jobs):
                self.entry_head(job, "title", "employer")
                self.description(job.get("description"), f"work_experience.{i}.description")

    def education(self):
        eds = self.cv.get("education") or []
        if eds:
            self.section_title(self.label("education"))
            for i, ed in enumerate(eds):
                self.entry_head(ed, "title", "organisation")
                for label, key in (("Field(s) of study", "field"), ("Final grade", "final_grade"), ("Thesis", "thesis")):
                    if ed.get(key):
                        self.labelled([(label, ed[key])])
                if ed.get("eqf_level"):
                    self.labelled([("Level in EQF", f"EQF level {ed['eqf_level']}")])
                if ed.get("description"):
                    self.description(ed["description"], f"education.{i}.description")
                else:
                    self.para(space_after=3)

    def languages(self):
        langs = self.cv.get("languages") or {}
        if not (langs.get("mother_tongues") or langs.get("other")):
            return
        self.section_title(self.label("languages"))
        if langs.get("mother_tongues"):
            self.labelled([("Mother tongue(s)", ", ".join(m.get("name", "") for m in langs["mother_tongues"]))])
        other = langs.get("other") or []
        if not other:
            return
        self.labelled([("Other language(s)", "")])
        if self.design["languages_layout"] == "list":
            keys = [("listening", "Listening"), ("reading", "Reading"),
                    ("spoken_production", "Spoken production"),
                    ("spoken_interaction", "Spoken interaction"), ("writing", "Writing")]
            for lang in other:
                p = self.para(space_after=1.5)
                self.run(p, lang.get("name", ""), bold=True, color=self.accent)
                levels = [f"{label} {lang[k]}" for k, label in keys if lang.get(k)]
                if levels:
                    self.run(p, "   " + " · ".join(levels), size=Pt(8 * self.design["scale"]))
            p = self.para(space_after=3)
            self.run(p, "Levels: A1 and A2: Basic user; B1 and B2: Independent user; C1 and C2: Proficient user",
                     italic=True, color=GREY_DATE, size=Pt(7.5))
            return
        table = self.doc.add_table(rows=2 + len(other), cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Mm(36)] + [Mm(30.4)] * 5
        heads = [("", 1), ("UNDERSTANDING", 2), ("SPEAKING", 2), ("WRITING", 1)]
        col = 0
        for text, span in heads:
            cell = table.cell(0, col)
            if span > 1:
                cell = cell.merge(table.cell(0, col + span - 1))
            self._cell(cell, text, bold=True, color=self.accent)
            col += span
        for i, text in enumerate(["", "Listening", "Reading", "Spoken production", "Spoken interaction", ""]):
            self._cell(table.cell(1, i), text, bold=bool(text))
        keys = ["listening", "reading", "spoken_production", "spoken_interaction", "writing"]
        for r, lang in enumerate(other, start=2):
            self._cell(table.cell(r, 0), lang.get("name", ""), bold=True, color=self.accent, left=True)
            for c, key in enumerate(keys, start=1):
                self._cell(table.cell(r, c), lang.get(key, ""))
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
                cell_border_bottom(cell)
        p = self.para(space_after=3)
        p.paragraph_format.space_before = Pt(3)
        self.run(p, "Levels: A1 and A2: Basic user; B1 and B2: Independent user; C1 and C2: Proficient user",
                 italic=True, color=GREY_DATE, size=Pt(7.5))

    def _cell(self, cell, text, bold=False, color=None, left=False):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
        if text:
            self.run(p, text, bold=bold, color=color, size=Pt(8 * self.design["scale"]))

    def digital_skills(self):
        skills = self.cv.get("digital_skills") or []
        if skills:
            self.section_title(self.label("digital_skills"))
            p = self.para()
            for i, s in enumerate(skills):
                if i:
                    self.run(p, " | ", color=GREY_SEP)
                self.run(p, s)

    def driving_licence(self):
        lic = self.cv.get("driving_licence") or []
        if lic:
            self.section_title(self.label("driving_licence"))
            self.labelled([("Driving licence", ", ".join(lic))])

    # --- typed optional sections -------------------------------------------
    def generic_section(self, key):
        """Nine sections share one entry shape; one writer covers them all."""
        items = self.cv.get(key) or []
        if not items:
            return
        self.section_title(self.label(key))
        for i, item in enumerate(items):
            if item.get("title"):
                p = self.para(space_after=0.5)
                self.run(p, item["title"], bold=True, color=self.accent)
            if item.get("organisation"):
                p = self.para(space_after=0.5)
                self.run(p, item["organisation"], italic=True)
            span = " – ".join(x for x in (item.get("start"), item.get("end") or
                                          ("Current" if item.get("start") else "")) if x)
            bits = [b for b in (f"[ {span} ]" if span else "",
                                f"[ {item['date']} ]" if item.get("date") else "",
                                item.get("location") or "") if b]
            if bits:
                p = self.para(space_after=0.5)
                self.run(p, "  ".join(bits), color=GREY_SEP)
            self.description(item.get("description"), f"{key}.{i}.description")
            pairs = [(k, v) for k, v in (("Mode of learning", item.get("mode")),
                                         ("Link", item.get("url"))) if v]
            if pairs:
                self.labelled(pairs)

    def publications(self):
        items = self.cv.get("publications") or []
        if not items:
            return
        self.section_title(self.label("publications"))
        for i, item in enumerate(items):
            if item.get("title"):
                p = self.para(space_after=0.5)
                self.run(p, item["title"], bold=True, color=self.accent)
            if item.get("year"):
                p = self.para(space_after=0.5)
                self.run(p, f"[ {item['year']} ]", color=GREY_SEP)
            self.description(item.get("description"), f"publications.{i}.description")
            pairs = [(k, v) for k, v in (("Reference", item.get("reference")),
                                         ("Authors", item.get("authors")),
                                         ("Journal Name", item.get("journal")),
                                         ("Volume, Issue and Pages", item.get("volume")),
                                         ("Publisher", item.get("publisher")),
                                         ("DOI", item.get("doi"))) if v]
            if pairs:
                self.labelled(pairs)

    def recommendations(self):
        items = self.cv.get("recommendations") or []
        if not items:
            return
        self.section_title(self.label("recommendations"))
        for i, item in enumerate(items):
            head = " | ".join(x for x in (item.get("name"), item.get("role")) if x)
            if head:
                p = self.para(space_after=0.5)
                self.run(p, head, bold=True, color=self.accent)
            self.description(item.get("description"), f"recommendations.{i}.description")
            phone = item.get("phone") or ""
            if phone and item.get("phone_dialing"):
                phone = f"(+{item['phone_dialing']}) {phone}"
            pairs = [(k, v) for k, v in (("Email", item.get("email")),
                                         ("Phone number", phone)) if v]
            if pairs:
                self.labelled(pairs)

    def custom(self, sec_id):
        for sec in self.cv.get("additional_sections") or []:
            if sec.get("id") == sec_id and sec.get("items"):
                si = (self.cv.get("additional_sections") or []).index(sec)
                self.section_title(sec.get("title", ""))
                for ii, item in enumerate(sec["items"]):
                    if item.get("heading"):
                        p = self.para(space_after=0.5)
                        self.run(p, item["heading"], bold=True, color=self.accent)
                    self.description(item.get("text"), f"additional_sections.{si}.items.{ii}.text")

    # --------------------------------------------------------------- main ---
    def build(self) -> bytes:
        self.header()
        for key in self.sections:
            if key.startswith("custom:"):
                self.custom(key.split(":", 1)[1])
            elif key in GENERIC_SECTIONS:
                self.generic_section(key)
            else:
                getattr(self, key)()
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()


def build_docx(cv: dict) -> bytes:
    return Writer(cv).build()
